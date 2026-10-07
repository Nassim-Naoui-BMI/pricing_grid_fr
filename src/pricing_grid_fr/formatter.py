from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from src.pricing_grid_fr import config


class ExcelFormatter:
    def __init__(self):
        self.font_base = Font(name="Century Gothic", size=9)
        self.font_bold = Font(name="Century Gothic", size=9, bold=True)
        self.font_white_bold = Font(
            name="Century Gothic", size=9, bold=True, color="FFFFFF"
        )

        # Couleurs issues du VBA
        self.fill_header = PatternFill(start_color="C0504D", fill_type="solid")  # Rouge
        self.fill_dark_header = PatternFill(
            start_color="BFBFBF", fill_type="solid"
        )  # Griz
        self.fill_gray = PatternFill(start_color="F2F2F2", fill_type="solid")
        self.fill_sub_header = PatternFill(start_color="DA8C9E", fill_type="solid")

        self.border_thin = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

    @staticmethod
    def autofit_columns(ws):
        """Ajuste automatiquement la largeur des colonnes en fonction de leur contenu."""
        for col in ws.columns:
            max_length = 0
            column_letter = col[0].column_letter

            for cell in col:
                try:
                    # On ignore les cellules vides
                    if cell.value:
                        # Si le texte contient des retours à la ligne, on prend la ligne la plus longue
                        cell_lines = str(cell.value).split("\n")
                        for line in cell_lines:
                            if len(line) > max_length:
                                max_length = len(line)
                except:
                    pass

            # On ajoute un coefficient multiplicateur ou une marge pour que le texte respire
            adjusted_width = max_length + 2

            # On applique la largeur calculée
            ws.column_dimensions[column_letter].width = adjusted_width

    def format_client_grid(self, filepath: str, is_derog: bool = False):
        wb = load_workbook(filepath)

        # ----------------------------------------------------------- Grille principale -----------------------------------------------------------
        ws = wb.active

        ws.views.sheetView[0].showGridLines = False  # Pas de quadrillage

        last_col = 14 if is_derog else 15  # Colonne N ou O[cite: 2]
        last_row = ws.max_row

        ws.auto_filter.ref = f"B7:{'N' if is_derog else 'O'}7"  # Appliquer le filtre

        # 1. En-tête gris et infos client[cite: 2]
        for row in ws.iter_rows(min_row=2, max_row=5, min_col=1, max_col=last_col):
            for cell in row:
                cell.fill = self.fill_gray
                cell.font = self.font_bold if cell.column == 1 else self.font_base
                cell.alignment = Alignment(horizontal="left")

        # 2. Ligne d'en-tête du tableau (Ligne 7)[cite: 2]
        for cell in ws[7][:last_col]:
            cell.fill = self.fill_header
            cell.font = self.font_white_bold
            cell.border = self.border_thin
            cell.alignment = Alignment(
                wrap_text=True, horizontal="center", vertical="center"
            )

        # 3. Corps du tableau et formats numériques[cite: 2]
        for row in ws.iter_rows(
            min_row=8, max_row=last_row, min_col=1, max_col=last_col
        ):
            for cell in row:
                cell.font = self.font_base
                cell.border = self.border_thin

                # Alignement centré pour les colonnes I à Fin[cite: 2]
                cell.alignment = Alignment(
                    horizontal="center" if cell.column >= 9 else "left",
                    vertical="center",
                )

                # Format Monétaire (Prix public, etc.)[cite: 2]
                if cell.column >= 12:
                    cell.number_format = "#,##0.00 €"

                # Format Pourcentage (Remise totale)[cite: 2]
                if not is_derog and cell.column == 13:
                    cell.number_format = "0.00%"

        # 4. Largeurs fixes des colonnes[cite: 2]
        widths = {
            "A": 14,
            "B": 12,
            "D": 40,
            "E": 22,
            "F": 22,
            "G": 33,
            "H": 33,
            "I": 9,
            "J": 9,
            "K": 10,
            "M": 9,
            "N": 12,
            "O": 9,
        }
        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        # ----------------------------------------------------------- Synthèse -----------------------------------------------------------
        if not is_derog:
            ws_synthese = wb["Synthèse"]

            first_level = [8, 24, 85]
            second_level = [9, 16, 25, 56, 86, 110]
            last_row = ws_synthese.max_row

            ws_synthese.views.sheetView[0].showGridLines = False  # Pas de quadrillage

            self.autofit_columns(ws_synthese)

            # 1. En-tête gris et infos client[cite: 2]
            for row in ws_synthese.iter_rows(
                min_row=2, max_row=5, min_col=1, max_col=2
            ):
                for cell in row:
                    cell.fill = self.fill_gray
                    cell.font = self.font_bold if cell.column == 1 else self.font_base
                    cell.alignment = Alignment(horizontal="left")

            ws_synthese.auto_filter.ref = "A7:B7"

            # 2. Ligne d'en-tête du tableau (Ligne 7)[cite: 2]
            for row in ws_synthese.iter_rows(
                min_row=7, max_row=7, min_col=1, max_col=2
            ):
                for cell in row:
                    cell.fill = self.fill_dark_header
                    cell.font = self.font_bold

            # 3. Tableau de synthèse
            for row in ws_synthese.iter_rows(
                min_row=7, max_row=last_row, min_col=1, max_col=2
            ):
                for cell in row:
                    if cell.row > 7:
                        cell.font = self.font_base
                    cell.border = self.border_thin
                    if cell.column == 2 and cell.row >= 10:
                        cell.number_format = "0.00%"
                        cell.alignment = Alignment(
                            horizontal="center", vertical="center"
                        )

            # 4. Niveau 1
            for row_number in first_level:
                for row in ws_synthese.iter_rows(
                    min_row=row_number, max_row=row_number, min_col=1, max_col=2
                ):
                    for cell in row:
                        cell.fill = self.fill_header
                        cell.font = self.font_white_bold

            # 5. Niveau 2
            for row_number in second_level:
                for row in ws_synthese.iter_rows(
                    min_row=row_number, max_row=row_number, min_col=1, max_col=2
                ):
                    for cell in row:
                        cell.fill = self.fill_sub_header
                        cell.font = (
                            self.font_bold if cell.column == 1 else self.font_base
                        )

        wb.save(filepath)
